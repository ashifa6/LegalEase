import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def _clean_inline_markdown(text: str) -> str:
    """
    Remove simple Markdown formatting.
    """

    text = text.replace("**", "")
    text = text.replace("__", "")
    text = text.replace("`", "")

    return text.strip()


def _add_normal_paragraph(
    document: Document,
    text: str,
):
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.15

    run = paragraph.add_run(
        _clean_inline_markdown(text)
    )

    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

    return paragraph


def _add_heading(
    document: Document,
    text: str,
    level: int,
):
    heading = document.add_heading(
        _clean_inline_markdown(text),
        level=level,
    )

    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(7)

    for run in heading.runs:

        run.font.name = "Times New Roman"

        if level == 1:
            run.font.size = Pt(14)
        else:
            run.font.size = Pt(12)

    return heading


def _add_bullet(
    document: Document,
    text: str,
):
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.left_indent = Inches(0.3)
    paragraph.paragraph_format.first_line_indent = Inches(-0.15)
    paragraph.paragraph_format.space_after = Pt(5)

    run = paragraph.add_run(
        "• " + _clean_inline_markdown(text)
    )

    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

    return paragraph


def _add_footer(document: Document):

    for section in document.sections:

        footer = section.footer

        paragraph = footer.paragraphs[0]

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(
            "LegalEase Inc. | contact@legalease.com | All Rights Reserved."
        )

        run.font.name = "Times New Roman"
        run.font.size = Pt(8)


def create_docx(document_text: str) -> bytes:
    """
    Create a professional legal DOCX document.
    """

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    # Default font.
    normal_style = document.styles["Normal"]

    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    # -----------------------------------------------------
    # LegalEase header
    # -----------------------------------------------------

    header = section.header

    header_paragraph = header.paragraphs[0]

    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    logo_run = header_paragraph.add_run(
        "⚖️ LegalEase"
    )

    logo_run.bold = True
    logo_run.font.name = "Times New Roman"
    logo_run.font.size = Pt(15)

    # -----------------------------------------------------
    # Clean document
    # -----------------------------------------------------

    document_text = document_text.replace(
        "\r\n",
        "\n",
    )

    document_text = document_text.replace(
        "\r",
        "\n",
    )

    document_text = re.sub(
        r"^```(?:markdown|md|text)?\s*",
        "",
        document_text,
        flags=re.IGNORECASE,
    )

    document_text = re.sub(
        r"\s*```$",
        "",
        document_text,
    )

    lines = document_text.split("\n")

    first_title_seen = False

    for line in lines:

        stripped = line.strip()

        if not stripped:
            continue

        # -------------------------------------------------
        # Main title
        # -------------------------------------------------

        if (
            stripped.startswith("# ")
            and not stripped.startswith("## ")
        ):

            title = stripped[2:].strip()

            paragraph = document.add_paragraph()

            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(18)

            run = paragraph.add_run(
                _clean_inline_markdown(title)
            )

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(18)

            first_title_seen = True

            continue

        # -------------------------------------------------
        # Level 2 heading
        # -------------------------------------------------

        if (
            stripped.startswith("## ")
            and not stripped.startswith("### ")
        ):

            heading = stripped[3:].strip()

            _add_heading(
                document,
                heading,
                level=1,
            )

            continue

        # -------------------------------------------------
        # Level 3 heading
        # -------------------------------------------------

        if stripped.startswith("### "):

            heading = stripped[4:].strip()

            _add_heading(
                document,
                heading,
                level=2,
            )

            continue

        # -------------------------------------------------
        # Bullet
        # -------------------------------------------------

        bullet_match = re.match(
            r"^[-*+]\s+(.*)",
            stripped,
        )

        if bullet_match:

            _add_bullet(
                document,
                bullet_match.group(1),
            )

            continue

        # -------------------------------------------------
        # Numbered clause
        # -------------------------------------------------

        numbered_match = re.match(
            r"^(\d+)[.)]\s+(.*)",
            stripped,
        )

        if numbered_match:

            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.line_spacing = 1.15

            run = paragraph.add_run(
                numbered_match.group(1)
                + ". "
                + _clean_inline_markdown(
                    numbered_match.group(2)
                )
            )

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

            continue

        # -------------------------------------------------
        # Lettered subclauses
        # -------------------------------------------------

        lettered_match = re.match(
            r"^([a-z])\)\s+(.*)",
            stripped,
            flags=re.IGNORECASE,
        )

        if lettered_match:

            paragraph = document.add_paragraph()

            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.first_line_indent = Inches(-0.15)
            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.paragraph_format.line_spacing = 1.1

            run = paragraph.add_run(
                lettered_match.group(1)
                + ") "
                + _clean_inline_markdown(
                    lettered_match.group(2)
                )
            )

            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

            continue

        # -------------------------------------------------
        # Signature fields
        # -------------------------------------------------

        if (
            stripped.startswith("By:")
            or stripped.startswith("Name:")
            or stripped.startswith("Title:")
            or stripped.startswith("Date:")
            or stripped.startswith("Signature:")
            or stripped.startswith("Authorized Representative:")
        ):

            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_after = Pt(5)

            run = paragraph.add_run(
                _clean_inline_markdown(
                    stripped
                )
            )

            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

            continue

        # -------------------------------------------------
        # Normal paragraph
        # -------------------------------------------------

        _add_normal_paragraph(
            document,
            stripped,
        )

    # -----------------------------------------------------
    # Footer
    # -----------------------------------------------------

    _add_footer(document)

    # -----------------------------------------------------
    # Save
    # -----------------------------------------------------

    output = BytesIO()

    document.save(output)

    return output.getvalue()